import asyncio
import os
import shutil
import tempfile
from pathlib import Path
import openpyxl

from config import AppConfig
from database import Database
from exporter import AvitoExporter
from models import AvitoItem, SearchQuery, SellerInfo
from parser_core import AvitoDataExtractor
from proxy_manager import ProxyManager


async def run_tests():
    print("=== 🧪 Запуск тестового набора Avito Parser ===")
    temp_dir = Path(tempfile.mkdtemp())
    try:
        # 1. Тест моделей
        print("1. Тестирование моделей Pydantic...")
        item = AvitoItem(
            id="123456789",
            title="iPhone 15 Pro 256GB Titanium",
            price=95000,
            price_string="95 000 ₽",
            url="https://www.avito.ru/moskva/telefony/iphone_15_pro_123456789",
            address="Москва, Тверская",
            metro="Охотный ряд",
            seller=SellerInfo(name="Александр", rating=4.95, reviews_count=42),
            params={"Встроенная память": "256 ГБ", "Цвет": "Титановый"},
            images=["https://10.img.avito.st/image/1/1280x960/sample1.jpg"]
        )
        assert item.id == "123456789"
        assert item.price == 95000
        print("   ✓ Модели работают корректно.")

        # 2. Тест базы данных SQLite
        print("2. Тестирование базы данных (aiosqlite)...")
        db = Database(db_path=temp_dir / "test.db")
        await db.init_db()

        # Сохранение нового товара
        saved, is_new, is_price_drop, prev = await db.save_item(item)
        assert is_new is True
        assert is_price_drop is False
        print("   ✓ Сохранение нового объявления: OK")

        # Обновление со снижением цены
        item_discount = item.model_copy()
        item_discount.price = 90000
        saved2, is_new2, is_price_drop2, prev2 = await db.save_item(item_discount)
        assert is_new2 is False
        assert is_price_drop2 is True
        assert prev2 == 95000
        print("   ✓ Фиксация снижения цены (95 000 -> 90 000 ₽): OK")

        # Поисковые запросы
        search_id = await db.add_search(SearchQuery(
            name="iPhone 15 Москва",
            url="https://www.avito.ru/moskva/telefony/iphone_15"
        ))
        searches = await db.get_searches()
        assert len(searches) == 1
        assert searches[0].id == search_id
        print("   ✓ Управление поисковыми запросами: OK")

        # Уведомления
        notif_sent_before = await db.is_notification_sent("123456789", search_id, "new")
        assert notif_sent_before is False
        await db.mark_notification_sent("123456789", search_id, "new")
        notif_sent_after = await db.is_notification_sent("123456789", search_id, "new")
        assert notif_sent_after is True
        print("   ✓ Дедупликация уведомлений: OK")

        # Статистика
        stats = await db.get_stats()
        assert stats["total_items"] == 1
        assert stats["active_searches"] == 1
        assert stats["total_price_drops"] == 1
        print("   ✓ Статистика БД: OK")

        # 3. Тест парсера данных и фоллбэка DOM
        print("3. Тестирование ядра извлечения данных (JSON & DOM)...")
        sample_html = """
        <!DOCTYPE html>
        <html>
        <head><title>Авито</title></head>
        <body>
            <div data-marker="item" data-item-id="998877">
                <a data-marker="item-title" href="/moskva/noutbuki/macbook_pro_998877">MacBook Pro M3 Max</a>
                <span data-marker="item-price">250 000 ₽</span>
                <span data-marker="item-address">Москва, метро Арбатская</span>
                <img data-marker="item-photo" src="//10.img.avito.st/1/test.jpg">
            </div>
        </body>
        </html>
        """
        extracted_dom = AvitoDataExtractor.extract_from_dom(sample_html)
        assert len(extracted_dom) == 1
        assert extracted_dom[0].id == "998877"
        assert extracted_dom[0].title == "MacBook Pro M3 Max"
        assert extracted_dom[0].price == 250000
        assert extracted_dom[0].address == "Москва, метро Арбатская"
        assert "test.jpg" in extracted_dom[0].main_image
        print("   ✓ Извлечение DOM через data-marker: OK")

        # Тест очистки цен
        assert AvitoDataExtractor.clean_price(" 1 250 000  ₽ ") == 1250000
        assert AvitoDataExtractor.clean_price("от 500 ₽/сут.") == 500
        assert AvitoDataExtractor.clean_price("Бесплатно") == 0
        print("   ✓ Хелпер очистки цен: OK")

        # 4. Тест экспорта в Excel, CSV, JSON
        print("4. Тестирование генерации отчетов...")
        exporter = AvitoExporter(export_dir=temp_dir)
        items_to_export = [item_discount, extracted_dom[0]]

        # Excel
        excel_file = exporter.export_to_excel(items_to_export)
        assert excel_file.exists()
        wb = openpyxl.load_workbook(excel_file)
        ws = wb.active
        assert ws.cell(row=1, column=2).value == "Заголовок"
        assert ws.cell(row=2, column=2).value == "iPhone 15 Pro 256GB Titanium"
        print("   ✓ Генерация Excel файла (.xlsx) со стилями и форматированием: OK")

        # CSV
        csv_file = exporter.export_to_csv(items_to_export)
        assert csv_file.exists()
        with open(csv_file, "r", encoding="utf-8-sig") as f:
            csv_content = f.read()
            assert "iPhone 15 Pro" in csv_content
            assert "MacBook Pro M3 Max" in csv_content
        print("   ✓ Генерация CSV файла (.csv) с UTF-8 BOM: OK")

        # JSON
        json_file = exporter.export_to_json(items_to_export)
        assert json_file.exists()
        print("   ✓ Генерация JSON файла (.json): OK")

        # 5. Тест менеджера прокси
        print("5. Тестирование менеджера прокси...")
        proxy_file = temp_dir / "proxies.txt"
        with open(proxy_file, "w", encoding="utf-8") as f:
            f.write("192.168.1.1:8080:user:pass\n")
            f.write("socks5://user2:pass2@10.0.0.1:1080\n")

        pm = ProxyManager(proxies_file=proxy_file)
        assert pm.total_count == 2
        p1 = pm.get_proxy()
        p2 = pm.get_proxy()
        assert "http://user:pass@192.168.1.1:8080" == p1
        assert "socks5://user2:pass2@10.0.0.1:1080" == p2
        print("   ✓ Парсинг форматов и ротация прокси: OK")

        print("\n[bold green]🎉 ВСЕ ТЕСТЫ УСПЕШНО ПРОЙДЕНЫ![/bold green]")

    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


if __name__ == "__main__":
    asyncio.run(run_tests())
